"""Editable DOCX export (REPORT-01). Mirrors the Markdown report; marks draft vs approved."""

import io

from docx import Document

from bionexus.evidence.eligibility import accepted_evidence
from bionexus.evidence.matrix import build_matrix
from bionexus.reporting.models import Report, ReportStatus
from bionexus.reporting.next_steps import propose_next_steps
from bionexus.workflow.state import RunState


def render_docx(state: RunState, report: Report) -> bytes:
    document = Document()
    document.add_heading(f"Evidence dossier: {state.request.disease}", level=1)
    document.add_paragraph(_status_line(report))
    document.add_paragraph(f"Question: {state.request.question}")
    document.add_paragraph(f"Evidence cutoff: {state.request.evidence_cutoff}")
    _add_matrix(document, state)
    _add_list(document, "Claims", [f"{c.text} [{', '.join(c.evidence_ids)}]" for c in state.claims])
    document.add_heading("Summary", level=2)
    document.add_paragraph(state.narrative)
    steps = propose_next_steps(state.issues, list(state.evidence.values()))
    _add_list(document, "Proposed next experiments", steps)
    limits = [i.description for i in state.issues]
    limits += [f"Excluded {e.evidence_id}: {e.reason}" for e in state.excluded]
    _add_list(document, "Limitations", limits)
    document.add_paragraph("Research use only. Not a treatment recommendation.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _status_line(report: Report) -> str:
    if report.status is ReportStatus.APPROVED:
        return f"Status: APPROVED by {report.approved_by} (version {report.version})"
    return f"Status: DRAFT, not scientist approved (version {report.version})"


def _add_matrix(document, state: RunState) -> None:
    document.add_heading("Evidence matrix", level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(
        table.rows[0].cells, ("Target", "Supporting", "Opposing", "No direction"), strict=True
    ):
        cell.text = text
    for row in build_matrix(accepted_evidence(list(state.evidence.values()))):
        values = (row.target, row.supporting_families, row.opposing_families, row.unknown_families)
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = str(value)


def _add_list(document, heading: str, items: list[str]) -> None:
    document.add_heading(heading, level=2)
    for item in items or ["None."]:
        document.add_paragraph(item, style="List Bullet")
